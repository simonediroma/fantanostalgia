from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from pydantic import BaseModel

from backend.api.db import get_db
from backend.api.routers.auth import get_current_admin

router = APIRouter(tags=["players"])

_ROLE_ALIASES = {"r", "ruolo", "ruo"}
_NAME_ALIASES = {"nome", "calciatore", "giocatore"}
_TEAM_ALIASES = {"squadra", "sq", "team"}
_QUOTA_ALIASES = {"qt a", "quotazione", "q.a.", "quota", "costo"}
_STARTS_ALIASES = {"pv", "presenze"}

_ROLE_MAP = {
    "p": "P", "por": "P",
    "d": "D", "dif": "D",
    "c": "C", "cen": "C",
    "a": "A", "att": "A",
}


def _find_columns(headers: list) -> dict:
    cols = {}
    for i, h in enumerate(headers):
        key = str(h).strip().lower() if h is not None else ""
        if key in _ROLE_ALIASES:
            cols.setdefault("role", i)
        elif key in _NAME_ALIASES:
            cols.setdefault("name", i)
        elif key in _TEAM_ALIASES:
            cols.setdefault("team", i)
        elif key in _QUOTA_ALIASES:
            cols.setdefault("quota", i)
        elif key in _STARTS_ALIASES:
            cols.setdefault("starts", i)
    for required in ("role", "name", "team", "quota"):
        if required not in cols:
            raise ValueError(f"Colonna obbligatoria mancante: {required}")
    return cols


def _safe_int(val, default: int = 0) -> int:
    try:
        return int(float(str(val))) if val not in (None, "") else default
    except (ValueError, TypeError):
        return default


def _is_rose_format(all_rows: list) -> bool:
    """Detect dual-column Rose format: role/name headers appear at both col 0 and col 5."""
    for row in all_rows:
        if len(row) < 7:
            continue
        c = [str(v).strip().lower() if v is not None else "" for v in row]
        if c[0] in _ROLE_ALIASES and c[1] in _NAME_ALIASES and c[5] in _ROLE_ALIASES and c[6] in _NAME_ALIASES:
            return True
    return False


def _parse_rose_rows(all_rows: list) -> tuple[list[dict], list[str]]:
    """Parse the dual-column Rose format (two fanta-teams side by side per block)."""
    rows_out = []
    skipped = 0
    left_team = None
    right_team = None

    for raw_row in all_rows:
        if all(v is None for v in raw_row):
            continue

        cells = [str(c).strip() if c is not None else "" for c in raw_row]

        # Team name row: col[0] has fanta-team name, col[1] empty, col[5] has right fanta-team name, col[6] empty
        if (cells[0]
                and not cells[1]
                and len(cells) > 5
                and cells[5]
                and (len(cells) < 7 or not cells[6])
                and cells[0].lower() not in _ROLE_ALIASES
                and not cells[0].lower().startswith("crediti")):
            left_team = cells[0]
            right_team = cells[5]
            continue

        # Skip header rows, metadata, crediti
        if (not left_team
                or cells[0].lower() in _ROLE_ALIASES
                or cells[0].lower().startswith("crediti")
                or cells[0].lower().startswith("*")
                or cells[0].lower().startswith("rose")
                or cells[0].lower().startswith("http")):
            continue

        # Left player
        role_left = _ROLE_MAP.get(cells[0].lower())
        name_left = cells[1] if len(cells) > 1 else ""
        if role_left and name_left:
            rows_out.append({
                "name": name_left,
                "role": role_left,
                "team": cells[2] if len(cells) > 2 else "",
                "quota": _safe_int(cells[3], default=1) or 1 if len(cells) > 3 else 1,
                "starts": 0,
                "fanta_team": left_team,
            })
        elif cells[0] and name_left:
            skipped += 1

        # Right player (cols 5-8)
        if len(cells) > 6 and cells[5]:
            role_right = _ROLE_MAP.get(cells[5].lower())
            name_right = cells[6] if len(cells) > 6 else ""
            if role_right and name_right:
                rows_out.append({
                    "name": name_right,
                    "role": role_right,
                    "team": cells[7] if len(cells) > 7 else "",
                    "quota": _safe_int(cells[8], default=1) or 1 if len(cells) > 8 else 1,
                    "starts": 0,
                    "fanta_team": right_team,
                })

    warnings = []
    if skipped:
        warnings.append(f"{skipped} righe saltate (ruolo non valido o nome vuoto)")
    return rows_out, warnings


def _parse_flat_rows(all_rows: list) -> tuple[list[dict], list[str]]:
    """Parse the standard flat format (single column block with header row)."""
    cols = None
    rows_out = []
    skipped = 0

    for raw_row in all_rows:
        if all(v is None for v in raw_row):
            continue

        cells = [str(c).strip() if c is not None else "" for c in raw_row]

        if cols is None:
            try:
                cols = _find_columns(cells)
            except ValueError:
                continue
            continue

        role_raw = cells[cols["role"]].lower()
        name = cells[cols["name"]]

        role = _ROLE_MAP.get(role_raw)
        if not role or not name:
            skipped += 1
            continue

        team = cells[cols["team"]]
        quota = _safe_int(cells[cols["quota"]], default=1) or 1

        starts = 0
        if "starts" in cols:
            starts = _safe_int(cells[cols["starts"]], default=0)

        rows_out.append({"name": name, "role": role, "team": team, "quota": quota, "starts": starts})

    if cols is None:
        raise ValueError("Header non trovato: colonne obbligatorie (ruolo, nome, squadra, quotazione) non trovate")

    warnings = []
    if "starts" not in cols:
        warnings.append("Colonna presenze non trovata, starts=0 per tutti")
    if skipped:
        warnings.append(f"{skipped} righe saltate (ruolo non valido o nome vuoto)")

    return rows_out, warnings


def _parse_excel(data: bytes) -> tuple[list[dict], list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    all_rows = list(ws.iter_rows(values_only=True))

    if _is_rose_format(all_rows):
        return _parse_rose_rows(all_rows)
    return _parse_flat_rows(all_rows)


def _require_league(conn, league_id: int):
    row = conn.execute("SELECT id FROM league WHERE id = ?", (league_id,)).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="Lega non trovata")


# ── Admin endpoints ─────────────────────────────────────────────────────────

@router.post("/admin/league/{league_id}/listone")
async def upload_listone(
    league_id: int,
    file: UploadFile,
    _: str = Depends(get_current_admin),
):
    data = await file.read()
    try:
        rows, warnings = _parse_excel(data)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception:
        raise HTTPException(status_code=400, detail="File Excel non valido o corrotto")

    by_role: dict[str, int] = {"P": 0, "D": 0, "C": 0, "A": 0}
    for r in rows:
        by_role[r["role"]] += 1

    teams_created: list[str] = []

    with get_db() as conn:
        _require_league(conn, league_id)

        # Build fanta-team → manager_id map for auto-assignment (Rose format)
        managers = conn.execute(
            "SELECT id, team_name FROM manager WHERE league_id = ?", (league_id,)
        ).fetchall()
        fanta_team_map: dict[str, int] = {
            m["team_name"].strip().lower(): m["id"] for m in managers
        }

        conn.execute("DELETE FROM player_current WHERE league_id = ?", (league_id,))
        for r in rows:
            manager_id = None
            fanta_team = r.get("fanta_team")
            if fanta_team:
                key = fanta_team.strip().lower()
                manager_id = fanta_team_map.get(key)
                if manager_id is None:
                    # Auto-create team from Excel; president can rename the manager later
                    cur = conn.execute(
                        "INSERT INTO manager (league_id, name, team_name) VALUES (?, ?, ?)",
                        (league_id, fanta_team.strip(), fanta_team.strip()),
                    )
                    manager_id = cur.lastrowid
                    fanta_team_map[key] = manager_id
                    teams_created.append(fanta_team.strip())
            conn.execute(
                "INSERT INTO player_current"
                " (league_id, name, role, team, quotation, starts_current_season, manager_id)"
                " VALUES (?, ?, ?, ?, ?, ?, ?)",
                (league_id, r["name"], r["role"], r["team"], r["quota"], r["starts"], manager_id),
            )

    return {
        "imported": len(rows),
        "by_role": by_role,
        "warnings": warnings,
        "teams_created": sorted(teams_created),
    }


class AssignItem(BaseModel):
    player_id: int
    manager_id: int


@router.post("/admin/league/{league_id}/assign")
def assign_players(
    league_id: int,
    body: list[AssignItem],
    _: str = Depends(get_current_admin),
):
    with get_db() as conn:
        _require_league(conn, league_id)

        player_ids = [item.player_id for item in body]
        manager_ids = list({item.manager_id for item in body})

        # Validate all players belong to this league
        placeholders = ",".join("?" * len(player_ids))
        found = conn.execute(
            f"SELECT id FROM player_current WHERE league_id = ? AND id IN ({placeholders})",
            [league_id, *player_ids],
        ).fetchall()
        if len(found) != len(player_ids):
            raise HTTPException(status_code=400, detail="Uno o più player_id non appartengono a questa lega")

        # Validate all managers belong to this league
        m_placeholders = ",".join("?" * len(manager_ids))
        found_mgr = conn.execute(
            f"SELECT id FROM manager WHERE league_id = ? AND id IN ({m_placeholders})",
            [league_id, *manager_ids],
        ).fetchall()
        if len(found_mgr) != len(manager_ids):
            raise HTTPException(status_code=400, detail="Uno o più manager_id non appartengono a questa lega")

        for item in body:
            conn.execute(
                "UPDATE player_current SET manager_id = ? WHERE id = ? AND league_id = ?",
                (item.manager_id, item.player_id, league_id),
            )

    return {"assigned": len(body)}


# ── Public endpoint ──────────────────────────────────────────────────────────

@router.get("/league/{league_id}/players")
def list_players(
    league_id: int,
    role: Optional[str] = None,
    manager_id: Optional[int] = None,
):
    with get_db() as conn:
        _require_league(conn, league_id)

        query = "SELECT * FROM player_current WHERE league_id = ?"
        params: list = [league_id]

        if role is not None:
            if role not in ("P", "D", "C", "A"):
                raise HTTPException(status_code=400, detail="Ruolo non valido. Valori: P, D, C, A")
            query += " AND role = ?"
            params.append(role)

        if manager_id is not None:
            query += " AND manager_id = ?"
            params.append(manager_id)

        query += " ORDER BY role, name"
        rows = conn.execute(query, params).fetchall()

    return [dict(r) for r in rows]
